from pathlib import Path
from copy import deepcopy
import json, csv
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

R=Path(__file__).resolve().parent.parent; W=R/'fuentes_y_referencias'; O=R
D=Document(W/'Plantilla_IEEE_A4_Original.docx')
sections=[deepcopy(s._sectPr) for s in D.sections]
body=D._element.body
for child in list(body):body.remove(child)
# La plantilla controla estilos, fuentes, geometría y dos columnas.
# Se omiten exclusivamente los bloques de autoría e instrucciones de ejemplo.
def p(text='',style='Body Text',bold=False):
    x=D.add_paragraph(style=style);r=x.add_run(text)
    if bold:r.bold=True
    if style in ['figure caption','table head','references']:
        n=OxmlElement('w:numPr');v=OxmlElement('w:numId');v.set(qn('w:val'),'0');n.append(v);x._p.get_or_add_pPr().append(n)
    return x
def h(text):
    x=p(text,'Heading 1');pr=x._p.get_or_add_pPr();n=OxmlElement('w:numPr');v=OxmlElement('w:numId');v.set(qn('w:val'),'0');n.append(v);pr.append(n)
    x.paragraph_format.keep_with_next=True
    return x
def col():
    x=p();x.paragraph_format.space_after=Pt(0);x.paragraph_format.space_before=Pt(0);x.paragraph_format.first_line_indent=Pt(0);x.paragraph_format.line_spacing=Pt(1);x.add_run().add_break(WD_BREAK.COLUMN)
def fig(name,caption,width=3.40):
    x=p();x.paragraph_format.first_line_indent=Pt(0);x.paragraph_format.space_before=Pt(4);x.paragraph_format.space_after=Pt(2);x.paragraph_format.keep_with_next=True;x.alignment=WD_ALIGN_PARAGRAPH.CENTER
    path=O/'resultados/figuras'/name
    x.add_run().add_picture(str(path),width=Inches(width))
    c=p(caption,'figure caption');c.paragraph_format.keep_with_next=False;c.paragraph_format.space_after=Pt(6)
def table(caption,headers,rows,widths):
    x=p(caption,'table head');x.paragraph_format.keep_with_next=True
    t=D.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','bottom','left','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'3');e.set(qn('w:color'),'B8B8B8');borders.append(e)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side,val in [('top',35),('bottom',35),('left',45),('right',45)]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    for j,(cell,text) in enumerate(zip(t.rows[0].cells,headers)):
        cell.width=Inches(widths[j]);x=cell.paragraphs[0];x.style=D.styles['table col head'];x.add_run(text).bold=True
        x.paragraph_format.first_line_indent=Pt(0);x.paragraph_format.space_after=Pt(0);x.paragraph_format.keep_with_next=True
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for j,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            x=cell.paragraphs[0];x.style=D.styles['table copy'];x.add_run(str(text));x.paragraph_format.first_line_indent=Pt(0);x.paragraph_format.space_after=Pt(0)
            x.paragraph_format.keep_with_next=False
        trpr=cells[0]._tc.getparent().get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    for j,width in enumerate(widths):t.columns[j].width=Inches(width)
    return t

p('Perfiles socioeducativos y predicción del rezago escolar mediante ENAHO y ESCALE 2025','paper title')
end=p();end.paragraph_format.space_after=Pt(0);end.paragraph_format.line_spacing=Pt(1);end.paragraph_format.space_before=Pt(0)
first=sections[0]
for tag in ['headerReference','footerReference']:
    for e in first.findall(qn('w:'+tag)):first.remove(e)
end._p.get_or_add_pPr().append(first)
main=sections[4]
for tag in ['headerReference','footerReference']:
    for e in main.findall(qn('w:'+tag)):main.remove(e)
typ=main.find(qn('w:type'))
if typ is None:typ=OxmlElement('w:type');main.insert(0,typ)
typ.set(qn('w:val'),'continuous')
body.append(main)

# Página 1, primera columna
p('Resumen—Se integraron microdatos del corte 2025 de ENAHO panel y registros de ESCALE para caracterizar perfiles socioeducativos y predecir un indicador operativo de rezago escolar. La cohorte clasificable comprende 19 490 personas de 6 a 17 años; 852 presentan desfase de dos o más años. Se reservaron 3 880 observaciones mediante conglomerados independientes y se compararon diez configuraciones de regresión logística, Random Forest, XGBoost y LightGBM con GroupKFold de tres particiones. K-Means con dos grupos obtuvo silueta 0,2681 y distinguió perfiles de capacidad económica y oferta educativa. XGBoost balanceado alcanzó ROC-AUC 0,6250, Average Precision 0,0921 y F1 positivo 0,1127 en prueba. SHAP identificó contribuciones asociadas con pobreza y contexto territorial. La baja precisión positiva y la calibración insuficiente restringen su utilidad a investigación exploratoria. La integración por área y la ponderación poblacional se examinan como complementos, sin atribuir causalidad ni representatividad nacional al clasificador.','Abstract')
p('Palabras clave—ENAHO, ESCALE, rezago escolar, validación por grupos, clustering, SHAP.','Keywords')
h('I. Introducción')
p('La brecha digital educativa combina recursos del hogar, acceso tecnológico y condiciones de la oferta escolar. Una asociación entre estas dimensiones y la trayectoria educativa puede orientar preguntas de investigación, pero no demuestra que una infraestructura determinada cause rezago. Además, los estudiantes de un mismo conglomerado comparten contexto; una partición individual aleatoria puede sobreestimar la capacidad de generalización.')
p('Este estudio integra dos fuentes oficiales de 2025 para responder qué perfiles socioeducativos aparecen en la muestra y cuánto permiten discriminar el rezago los atributos observados. Se utiliza ENAHO para condiciones individuales y del hogar [1], y ESCALE para indicadores contextuales de oferta educativa [2]. La unidad analítica es la persona; los indicadores distritales no identifican su escuela de matrícula.')
p('La contribución consiste en una comparación reproducible de arquitecturas, una partición por conglomerado y una interpretación explícita de sus errores. La evaluación privilegia la clase minoritaria y la curva Precision–Recall, porque una alta exactitud puede coexistir con la omisión de todos los casos positivos.')
col()

# Página 1, segunda columna
h('II. Metodología de datos')
p('Se empleó el corte identificado con sufijo _25 de los archivos panel 2021–2025 de ENAHO: Vivienda y Hogar, Miembros, Educación y Sumaria. No se presume equivalencia demostrada con la base anual independiente. Los archivos de origen contienen 44 599 registros de vivienda, 115 145 de personas, 106 958 de educación y 33 702 de Sumaria en los extractos utilizados.')
p('Se preservaron identificadores como texto y sus ceros iniciales. Las fuentes CSV se leyeron con la codificación correspondiente y se normalizaron blancos y espacios. La unión personal utiliza conglomerado, vivienda, hogar y persona del corte 2025; la unión con Vivienda y Sumaria omite persona. Se validaron unicidad de las tablas de origen y cardinalidades uno a uno o muchos a uno según el nivel de registro.')
fig('integracion-enaho-escale.drawio.png','Fig. 1. Integración de módulos y enlace contextual distrital. El complemento por área utiliza además el área de residencia.',3.40)
p('El universo de 6 a 17 años contiene 24 195 personas. El enlace con Educación encuentra 23 668 coincidencias y 527 ausencias. Se conservó un maestro para auditar estas pérdidas; el objetivo no se imputó. Tras exigir matrícula y grado válido en primaria o secundaria se obtuvieron 19 490 registros clasificables y 4 705 exclusiones respecto de la población inicial.')
p('ESCALE se agregó por CODGEO, homologado al UBIGEO de seis dígitos. El ratio distrital se calculó como suma de alumnos dividida entre suma de docentes; una media de ratios por escuela daría otra medida. Se incluyeron proporción de servicios públicos y número de locales activos. Estas variables describen el entorno y no recursos individuales disponibles.')
col()

# Página 2, primera columna
h('III. Limpieza y construcción del objetivo')
p('La limpieza se concentró en las variables y uniones del problema analítico. Se comprobaron llaves completas, duplicados, conversiones numéricas, faltantes y rangos. Los indicadores de cruce se retuvieron para trazabilidad. Los valores ausentes de los predictores se imputaron con mediana o moda dentro del entrenamiento de cada partición; la clasificación no incorporó personas con objetivo indeterminado.')
p('La edad normativa operativa es grado + 5 para primaria y año + 11 para secundaria. El rezago toma valor uno si la edad observada supera esa referencia en dos o más años. Se utilizó P308C para grado de primaria y P308B para año de secundaria; interpretar el cero estructural de P308B como primer grado produciría etiquetas incorrectas.')
table('TABLA I. CONSTRUCCIÓN DEL UNIVERSO', ['Condición','Registros'],[['Personas de 6 a 17 años','24 195'],['Con enlace educativo','23 668'],['Con objetivo clasificable','19 490'],['Sin rezago / con rezago','18 638 / 852'],['Excluidos del universo inicial','4 705']],[2.25,1.15])
p('El umbral de dos años es consistente con la definición de atraso de ESCALE [3], pero la edad de ENAHO corresponde a la entrevista y la referencia administrativa utiliza el 31 de marzo. Por tanto, las etiquetas son una aproximación y no una reproducción exacta del indicador SIAGIE. La selección no permite certificar el aislamiento completo de modalidades EBR y CEBA, porque no se incorporó P307D.')
p('En P1121 y P1144, el código 1 representa la alternativa declarada y 0 un pase. Se conservaron sus categorías originales; el cero no se interpretó como una ausencia de servicio confirmada en todos los hogares. Las categorías especiales de lengua materna permanecen separadas. Esta ambigüedad limita las conclusiones sobre brecha digital.')
p('Los extremos económicos no se eliminaron automáticamente. En entrenamiento, 1,74 % de ingresos y 1,45 % de gastos superaron Q3 + 3·IQR. No hubo fallos de conversión, infinitos ni valores negativos en los predictores numéricos inspeccionados. Un valor alto plausible no constituye por sí solo un error. En clustering se aplicó log1p a ingreso, gasto y locales; los montos son totales anuales del hogar, no per cápita.')
col()

# Página 2, segunda columna
h('IV. Diseño experimental y control de fuga')
p('GroupShuffleSplit con semilla 42 reservó conglomerados completos para prueba: 15 610 personas de entrenamiento y 3 880 de prueba, con 701 y 151 positivos, respectivamente. La intersección de conglomerados fue vacía. Dentro del entrenamiento, GroupKFold de tres particiones [4] produjo validaciones de 5 204, 5 203 y 5 203 personas; se verificó otra vez la separación de grupos en cada fold.')
p('Cada pipeline aprendió mediana, moda, StandardScaler y One-Hot Encoding exclusivamente del entrenamiento del fold. Las categorías desconocidas se toleraron sin volver a ajustar el codificador. El sobremuestreo aleatorio se aplicó únicamente al entrenamiento de regresión logística y Random Forest. Los pesos de clase y scale_pos_weight se calcularon con el entrenamiento correspondiente. La validación conservó su prevalencia original.')
p('Se excluyeron edad, grado, edad normativa, desfase y variables de construcción del objetivo del clasificador; tampoco ingresaron llaves ni indicadores de correspondencia. Permanecieron once predictores: sexo, lengua, códigos de electricidad e internet, pobreza, dominio, ingreso, gasto, ratio docente, proporción pública y locales. Separar conglomerados reduce dependencia intragrupo, pero no garantiza distritos completos inéditos ni validación temporal.')
p('Se comparó regresión logística con Random Forest [5], XGBoost [6] y LightGBM [7]. Se mantuvieron diez configuraciones: normal, balanceada y sobremuestreada para los dos primeros; normal y balanceada para los dos últimos. Random Forest utilizó 120 árboles y profundidad 12; XGBoost, 150 árboles, profundidad 4 y tasa 0,05; LightGBM, 150 árboles y tasa 0,05. El notebook conserva los parámetros restantes.')
p('La selección histórica utilizó el F1 de rezago medio en GroupKFold. El umbral final permaneció en 0,5. Como sensibilidad posterior limitada, se evaluó en entrenamiento una cuadrícula XGBoost de profundidad {3, 4} y árboles {100, 150}; no se consultó prueba para elegir una nueva configuración ni se reemplazó el modelo histórico.')
table('TABLA II. SENSIBILIDAD ACOTADA EN TRES FOLDS', ['Prof.','Árboles','F1 positivo','DE'],[['3','100','0,12254','0,00784'],['3','150','0,12216','0,00882'],['4','100','0,12023','0,00680'],['4','150','0,12030','0,00955']],[.48,.61,1.25,1.06])
col()

# Página 3, primera columna
h('V. Exploración y ponderación descriptiva')
p('El rezago representa 4,37 % de la cohorte clasificable. La prevalencia es 4,49 % en entrenamiento y 3,89 % en prueba; la diferencia se conserva porque la separación sigue conglomerados y no una estratificación individual. El EDA examinó edades, distribución de clases, pobreza, dominio, relaciones ingreso–gasto, ausentes y valores extremos. Sus figuras se calcularon sobre entrenamiento para no orientar la selección con información de prueba.')
fig('articulo/eda_pobreza.png','Fig. 2. Rezago por condición de pobreza en entrenamiento. Son proporciones muestrales, no efectos causales.')
p('La mayor tasa observada en hogares pobres motiva examinar condiciones socioeconómicas junto con la infraestructura contextual. La correlación entre ingreso y gasto puede repartir la atribución predictiva entre variables; no equivale a dos causas independientes. El número de locales distritales también refleja tamaño y concentración territorial, no acceso efectivo de cada estudiante.')
p('Se distinguió FACTOR07 de Sumaria del factor personal FACPOB07. Los modelos y las métricas predictivas originales no se ponderaron. Un complemento utilizó FACPOB07 positivo y finito entre residentes habituales de entrenamiento; la condición de residencia se construyó con P204, P205 y P206 según la metodología INEI [8]. Los pesos poblacionales no se confundieron con el balanceo de clases.')
table('TABLA III. REZAGO EN RESIDENTES DE ENTRENAMIENTO', ['Dominio','n','Muestral %','Ponderado %'],[['Total','15 448','4,4731','4,1602'],['Urbano','9 828','3,8970','3,8890'],['Rural','5 620','5,4804','5,1964']],[.73,.67,1.0,1.0])
p('Estas tasas describen el dominio analítico de entrenamiento. La suma de sus pesos no se interpreta como población nacional, pues no se reajustó por la partición ni se calcularon intervalos con estratos y conglomerados del diseño complejo. La diferencia rural–urbana es descriptiva; no identifica un efecto del área.')
col()

# Página 3, segunda columna
h('VI. Resultados de clustering')
p('Se compararon K-Means y mezclas gaussianas para k de 2 a 6, tras StandardScaler y sobre entrenamiento. La versión refinada excluyó P1121/P1144 para evitar perfiles dominados por sus códigos. Se utilizaron cinco variables continuas y desfase escolar. K-Means con k = 2 maximizó silueta y Calinski–Harabasz y presentó el menor Davies–Bouldin entre las alternativas examinadas.')
table('TABLA IV. COMPARACIÓN NO SUPERVISADA', ['Método','k','Silueta','DB','CH'],[['K-Means','2','0,2681','1,4293','6 725'],['K-Means','3','0,1983','1,7108','4 908'],['K-Means','4','0,1869','1,6118','4 293'],['K-Means','5','0,1833','1,4886','3 934'],['K-Means','6','0,1826','1,5040','3 723'],['GMM','2','0,1836','1,6762','3 896'],['GMM','3','0,1426','2,1887','3 916'],['GMM','4','0,1204','2,2091','3 289'],['GMM','5','0,1198','3,8064','1 547'],['GMM','6','0,1099','5,1331','1 300']],[.79,.25,.8,.79,.77])
fig('articulo/clustering.png','Fig. 3. PCA y silueta del K-Means refinado. Las dos componentes explican 63,76 % de la varianza; la silueta se visualiza en una submuestra.',3.15)
table('TABLA V. CARACTERIZACIÓN DE PERFILES', ['Indicador','Grupo 0','Grupo 1'],[['Personas','8 310','7 300'],['Ingreso mediano S/ anual','47 505,88','22 591,46'],['Gasto mediano S/ anual','40 020,64','21 410,77'],['Ratio alumnos/docente','17,67','12,00'],['Proporción pública','0,59','1,00'],['Locales escolares','126','44'],['Desfase medio / rezago %','0,38 / 3,84','0,45 / 5,23']],[1.6,.9,.9])
col()

# Página 4, primera columna
h('VII. Comparación supervisada')
p('El grupo 0 se denomina «mayor capacidad económica y contexto mixto» y el grupo 1 «menor capacidad y predominio público». Su separación es moderada. El desfase intervino en la construcción de los grupos, por lo que su diferencia en rezago no es una validación independiente. No se infiere aislamiento digital ni ruralidad a partir de estas etiquetas.')
table('TABLA VI. MEDIAS DE VALIDACIÓN GROUPKFOLD', ['Configuración','F1 macro','F1 +','Recall +','ROC'],[['RL normal','0,4885','0,0000','0,0000','0,6351'],['RL pesos','0,4405','0,1166','0,5461','0,6340'],['RL sobre-muestreo','0,4405','0,1151','0,5351','0,6329'],['RF normal','0,4885','0,0000','0,0000','0,6273'],['RF pesos','0,5158','0,1155','0,2232','0,6070'],['RF sobre-muestreo','0,5171','0,1040','0,1675','0,5851'],['XGB normal','0,4885','0,0000','0,0000','0,6182'],['XGB pesos','0,4809','0,1203','0,4081','0,6152'],['LGB normal','0,4885','0,0000','0,0000','0,6076'],['LGB pesos','0,4885','0,1133','0,3377','0,6054']],[1.16,.6,.51,.62,.51])
p('RL: regresión logística; RF: Random Forest; XGB: XGBoost; LGB: LightGBM. «+» representa rezago. El F1 macro más alto correspondió a RF sobremuestreado; XGBoost balanceado fue seleccionado por F1 positivo, no por superioridad en todas las métricas. Las variantes normales no recuperaron positivos al umbral fijado. El balanceo elevó sensibilidad a costa de precisión.')
fig('articulo/roc_pr.png','Fig. 4. Evaluación final de XGBoost balanceado. La referencia PR horizontal corresponde a la prevalencia de prueba de 3,89 %.')
col()

# Página 4, segunda columna
h('VIII. Evaluación final y análisis de errores')
table('TABLA VII. PRUEBA INDEPENDIENTE POR CONGLOMERADO', ['Métrica','Valor'],[['Exactitud','0,736082'],['Precisión de rezago','0,064806'],['Recall de rezago','0,430464'],['F1 de rezago','0,112652'],['F1 macro','0,478821'],['ROC-AUC','0,625042'],['Average Precision','0,092119'],['PR-AUC trapezoidal','0,089826'],['Brier','0,188955']],[2.17,1.23])
fig('articulo/errores_calibracion.png','Fig. 5. Matriz de confusión y calibración en prueba. La diagonal representa calibración ideal, no una meta alcanzada.')
p('De 151 positivos, el modelo detectó 65 y omitió 86: la tasa de falsos negativos es 56,95 %. Entre 1 003 alertas, 938 son falsos positivos; solo 6,48 % corresponden a rezago observado. Un uso operativo generaría una carga elevada de verificaciones y dejaría sin detectar a más de la mitad de los casos. Las alertas no justifican decisiones automáticas sobre estudiantes.')
p('El ROC-AUC de 0,6250 muestra discriminación limitada. Average Precision supera la prevalencia de prueba, pero mantiene un nivel absoluto bajo; no es idéntica al área trapezoidal PR. La curva de calibración evidencia probabilidades sobreestimadas en varios intervalos. El Brier 0,1890 es peor que un predictor constante próximo a la prevalencia de entrenamiento (aproximadamente 0,0374 en esta prueba). La exactitud no demuestra utilidad ni superioridad frente a predecir siempre la mayoría.')
col()

# Página 5, primera columna
h('IX. Explicabilidad global y dependencias')
p('Se calcularon contribuciones TreeSHAP del XGBoost final sobre 300 filas de entrenamiento seleccionadas con semilla 42. SHAP descompone el resultado del modelo en contribuciones aditivas [9]; aquí se expresa en log-odds. La figura muestra diez características individuales y una fila que agrupa las restantes. Las categorías One-Hot son características separadas, no diez variables originales distintas.')
fig('articulo/shap_global.png','Fig. 6. Beeswarm SHAP de entrenamiento. Un valor positivo desplaza la predicción hacia rezago; el color representa el valor de la característica.',3.20)
p('La categoría no pobre tiende a disminuir la salida frente a su ausencia; Sierra sur y Sierra centro muestran contribuciones negativas en numerosas observaciones. Estos resultados dependen de la muestra, la codificación y las demás variables. La presencia del indicador castellano no demuestra un efecto causal de la lengua ni permite calificar el desempeño de todas las lenguas originarias.')
fig('articulo/shap_dependencias.png','Fig. 7. Dependencias de alumnos/docente y proporción pública. El color representa el otro indicador contextual.',3.05)
col()

# Página 5, segunda columna
p('Las dependencias no muestran una relación monótona simple. Ratios similares producen contribuciones distintas según el contexto; en proporciones públicas cercanas a uno aparecen contribuciones positivas y negativas. El color ayuda a explorar heterogeneidad, pero no constituye una prueba formal de interacción causal. Se trata de ratios distritales compartidos por varias personas.')
h('X. Explicaciones individuales')
fig('articulo/waterfall_positivo.png','Fig. 8. Verdadero positivo de probabilidad más alta entre los casos positivos correctamente clasificados de prueba.')
p('El caso positivo alcanza probabilidad 0,8592 y salida aproximada de 1,809 log-odds. Las barras rojas aumentan la salida respecto del valor base y las azules la reducen. La suma de contribuciones reproduce la predicción del modelo; no representa una probabilidad independiente por variable. Su selección es ilustrativa y no estima el rendimiento de una población.')
fig('articulo/waterfall_negativo.png','Fig. 9. Verdadero negativo de probabilidad más baja entre los casos negativos correctamente clasificados de prueba.')
p('El caso negativo alcanza probabilidad 0,0402 y salida aproximada de −3,174 log-odds. Las contribuciones protectoras del contexto y la condición socioeconómica compensan las que elevan la salida. Los dos casos corresponden a observaciones distintas y no a una intervención antes–después. Mostrar aciertos extremos no sustituye el examen de los 938 falsos positivos y 86 falsos negativos.')
p('La atribución entre variables correlacionadas puede cambiar al modificar el conjunto de predictores. La importancia global tampoco certifica equidad: se requieren errores por lengua y área con tamaños suficientes e incertidumbre.')
col()

# Página 6, primera columna
h('XI. Discusión y alcance de los resultados')
p('El contraste entre arquitecturas muestra que la disponibilidad de modelos complejos no resuelve la escasez de señal predictiva. El pequeño aumento del F1 positivo en la sensibilidad de profundidad 3 es menor que la desviación entre folds; no establece una mejora concluyente. Se mantuvo el modelo histórico para evitar una nueva selección apoyada en la prueba ya observada. Una validación futura requeriría otro periodo o conjunto externo.')
p('La integración distrital permite relacionar hogares con oferta escolar, pero combina estudiantes con contextos heterogéneos. Como complemento se agregaron servicios activos de primaria B0 y secundaria F0 por distrito y área censal. Se preservó la unidad COD_MOD–ANEXO. ESTRATO 1–5 se mapeó a urbano y 6–8 a rural; en ESCALE se utilizó AREA_CENSO corroborada con DAREACENSO.')
table('TABLA VIII. COBERTURA DEL COMPLEMENTO POR ÁREA', ['Área','Registros','Enlace','Sin enlace'],[['Urbana','12 478','12 152','326'],['Rural','7 012','6 961','51'],['Total','19 490','19 113','377']],[.68,.89,.88,.95])
p('Las 377 observaciones sin correspondencia permanecen ausentes en ese complemento. No se rellenaron con promedios distritales ni se sustituyeron los predictores históricos. La comparación original conserva su alcance distrital; el complemento satisface la desagregación contextual, pero no mide asistencia a una escuela concreta. La cobertura desigual debe evaluarse antes de ampliar el modelo.')
p('La definición operativa del objetivo introduce incertidumbre por fecha de edad y modalidad educativa. También existe una diferencia entre el panel disponible y la fuente anual indicada como referencia. No se dispone de evidencia para convertir los resultados en estimaciones representativas de todos los escolares del Perú. Las tasas ponderadas corresponden solo al dominio de residentes analizado.')
p('El clustering describe gradientes de recursos económicos y oferta educativa, con superposición apreciable. Las variables refinadas no incluyen competencias digitales, frecuencia de uso TIC, altitud ni conectividad escolar directa. Por ello, los perfiles no pueden presentarse como tipologías completas de exclusión digital. Incorporar estos campos exige verificar su disponibilidad y semántica y realizar una nueva evaluación independiente.')
p('Para política educativa, los resultados respaldan focalizar la investigación y la verificación territorial de carencias, no adjudicar beneficios mediante puntuaciones individuales. Conviene contrastar oferta docente, conectividad efectiva y condiciones económicas con diagnóstico local, y evaluar intervenciones mediante diseños específicos. La asociación de pobreza y rezago orienta hipótesis; no cuantifica el impacto de una transferencia ni de instalar internet.')
col()

# Página 6, segunda columna
h('XII. Conclusiones')
p('La integración produjo una cohorte trazable de 19 490 personas y una comparación reproducible con separación estricta por conglomerado. K-Means con dos grupos obtuvo la mejor calidad relativa entre los métodos explorados, aunque su silueta indica separación moderada. Los perfiles reflejan diferencias económicas y de oferta educativa sin demostrar aislamiento digital.')
p('XGBoost balanceado fue el mejor de las diez configuraciones según F1 positivo de validación, pero su evaluación final revela baja precisión y calibración insuficiente. SHAP permite examinar contribuciones globales y casos particulares sin convertirlas en explicaciones causales. La evidencia apoya continuar la validación y mejorar la medición de brecha digital; no respalda decisiones automatizadas sobre estudiantes.')
p('El código, el diccionario, los datos procesados, las tablas y las figuras permiten auditar el procedimiento. La reproducción desde kernel limpio recuperó las métricas y perfiles guardados. Una ampliación deberá armonizar el universo anual y la edad de referencia, integrar contexto por área en un nuevo experimento y reservar datos independientes para evaluar cualquier mejora.')
h('Referencias')
refs=[
'[1] INEI, “ENAHO. Ficha técnica 2025 y diccionario de datos panel 2021–2025”, módulos 100, 200, 300 y Sumaria, 2025. Documentación de las fuentes empleadas.',
'[2] MINEDU, “ESCALE. Padrón de servicios educativos y locales escolares 2025”, archivos 00_Padron y 00_Padlocal, 2025. Disponible: https://escale.minedu.gob.pe/',
'[3] MINEDU, “Porcentaje de alumnos con atraso escolar”, ficha metodológica de indicadores ESCALE. Disponible: escale.minedu.gob.pe/tendencias-2016-portlet/servlet/tendencias/archivo?idCuadro=14&tipo=meta',
'[4] Scikit-learn developers, “GroupKFold”, documentación oficial. Disponible: https://scikit-learn.org/stable/modules/generated/sklearn.model_selection.GroupKFold.html (consulta: oct. 2026).',
'[5] L. Breiman, “Random forests”, Machine Learning, vol. 45, pp. 5–32, 2001, doi: 10.1023/A:1010933404324.',
'[6] T. Chen y C. Guestrin, “XGBoost: A scalable tree boosting system”, Proc. ACM SIGKDD, pp. 785–794, 2016, doi: 10.1145/2939672.2939785.',
'[7] G. Ke et al., “LightGBM: A highly efficient gradient boosting decision tree”, Advances in Neural Information Processing Systems, vol. 30, 2017.',
'[8] INEI, “Metodología de cálculo de los factores de expansión”, documentación metodológica ENAHO. Disponible: https://www.inei.gob.pe/media/MenuRecursivo/publicaciones_digitales/Est/Lib1244/Libro.pdf',
'[9] S. M. Lundberg y S.-I. Lee, “A unified approach to interpreting model predictions”, Advances in Neural Information Processing Systems, vol. 30, 2017.'
]
for text in refs:
    x=p(text,'references');x.paragraph_format.keep_with_next=False;x.alignment=WD_ALIGN_PARAGRAPH.LEFT

# Metadatos personales vacíos y sin material editorial de la plantilla.
D.core_properties.author='';D.core_properties.last_modified_by='';D.core_properties.title='Perfiles socioeducativos y predicción del rezago escolar mediante ENAHO y ESCALE 2025';D.core_properties.subject='Integración ENAHO y ESCALE 2025';D.core_properties.comments=''
for sec in D.sections:
    for part in [sec.header,sec.footer]:
        for x in part.paragraphs:
            x.clear()
D.save(O/'Articulo_Cientifico_IA.docx')
(O/'fuentes_y_referencias/Referencias_Articulo.txt').write_text('\n\n'.join(refs),encoding='utf-8')
print('DOCX',O/'Articulo_Cientifico_IA.docx','PARAGRAPHS',len(D.paragraphs),'FIGURES',len(D.inline_shapes))
